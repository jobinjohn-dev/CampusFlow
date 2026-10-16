"""Build the CampusFlow Phase 3 final project report from verified artifacts."""

from __future__ import annotations

from datetime import date
from pathlib import Path
import re
import subprocess
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
PHASE3 = ROOT / "phase3"
OUTPUT = ROOT / "deliverables" / "phase3" / "CampusFlow_Phase3_Final_Project_Report.docx"

NAVY = "17365D"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"
CODE_GRAY = "F3F4F6"
BLACK = RGBColor(0, 0, 0)


def require_tree() -> dict[str, str | int]:
    modules = [
        "model.py", "lexer.py", "ast_nodes.py", "parser.py", "symbols.py",
        "semantic.py", "ir.py", "optimizer.py", "interpreter.py",
        "metrics.py", "pipeline.py", "cli.py",
    ]
    examples = [
        "valid.cflow", "optimization_demo.cflow", "boundary_cases.cflow",
        "complex_schedule.cflow", "invalid.cflow", "syntax_errors.cflow",
        "semantic_errors.cflow", "resource_conflict.cflow",
    ]
    for name in modules:
        if not (PHASE3 / "campusflow" / name).is_file():
            raise FileNotFoundError(f"Required module is missing: {name}")
    for name in examples:
        if not (PHASE3 / "examples" / name).is_file():
            raise FileNotFoundError(f"Required example is missing: {name}")

    tests = subprocess.run(
        [sys.executable, "-m", "unittest", "discover", "-s", "tests", "-v"],
        cwd=PHASE3,
        text=True,
        capture_output=True,
        check=False,
    )
    test_text = tests.stdout + tests.stderr
    if tests.returncode:
        raise RuntimeError("The Phase 3 test suite failed before report generation")
    test_match = re.search(r"Ran (\d+) tests", test_text)
    if not test_match:
        raise RuntimeError("Could not determine the verified test count")

    optimizer = subprocess.run(
        [sys.executable, "run.py", "examples/optimization_demo.cflow", "--stage", "optimized"],
        cwd=PHASE3,
        text=True,
        capture_output=True,
        check=True,
    ).stdout
    benchmark = (PHASE3 / "evidence" / "benchmark_output.txt").read_text(encoding="utf-8")

    def value(text: str, label: str) -> str:
        match = re.search(rf"^{re.escape(label)}:\s*(.+)$", text, re.MULTILINE)
        if not match:
            raise RuntimeError(f"Missing measured field: {label}")
        return match.group(1).strip()

    return {
        "test_count": int(test_match.group(1)),
        "demo_original": value(optimizer, "Original instructions"),
        "demo_optimized": value(optimizer, "Optimized instructions"),
        "demo_reduction": value(optimizer, "Instruction reduction"),
        "bench_iterations": value(benchmark, "Benchmark iterations"),
        "bench_original": value(benchmark, "Original instruction count"),
        "bench_optimized": value(benchmark, "Optimized instruction count"),
        "bench_reduction": value(benchmark, "Instruction reduction"),
        "original_ns": value(benchmark, "Original median execution time (ns)"),
        "optimized_ns": value(benchmark, "Optimized median execution time (ns)"),
        "ratio": value(benchmark, "Measured time ratio (original/optimized)"),
        "original_bytes": value(benchmark, "Original peak memory (bytes)"),
        "optimized_bytes": value(benchmark, "Optimized peak memory (bytes)"),
        "equivalence": value(benchmark, "Execution equivalence"),
        "platform": value(benchmark, "Platform"),
        "python": value(benchmark, "Python"),
    }


def set_cell_fill(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), LIGHT_GRAY)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    header_properties = header._tr.get_or_add_trPr()
    header_properties.append(OxmlElement("w:tblHeader"))
    header_properties.append(OxmlElement("w:cantSplit"))
    for index, text in enumerate(headers):
        cell = header.cells[index]
        set_cell_fill(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        cells = row.cells
        for index, text in enumerate(values):
            cell = cells[index]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_fill(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            run = paragraph.add_run(str(text))
            run.font.size = Pt(8.7)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_code(doc: Document, text: str, size: float = 8.5) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.2)
    paragraph.paragraph_format.right_indent = Inches(0.2)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    if len(text.splitlines()) <= 25:
        paragraph.paragraph_format.keep_together = True
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), CODE_GRAY)
    paragraph._p.get_or_add_pPr().append(shading)
    run = paragraph.add_run(text.rstrip())
    run.font.name = "Consolas"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
    run.font.size = Pt(size)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(item, style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_after = Pt(14)
    title_properties = title._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 17, 14, 7),
        ("Heading 2", 13, 11, 5),
        ("Heading 3", 11, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for section in doc.sections:
        section.different_first_page_header_footer = True
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header_run = header.add_run("CampusFlow Phase 3 Final Project Report")
        header_run.font.size = Pt(8)
        header_run.font.color.rgb = BLACK
        add_page_number(section.footer.paragraphs[0])


def extract_block(path: Path, start: str, end: str, max_lines: int = 48) -> str:
    lines = path.read_text(encoding="utf-8").splitlines()
    first = next(index for index, line in enumerate(lines) if start in line)
    last = next((index for index in range(first + 1, len(lines)) if end in lines[index]), len(lines))
    return "\n".join(lines[first:min(last, first + max_lines)])


def build_report(facts: dict[str, str | int]) -> None:
    doc = Document()
    configure_styles(doc)

    cover = doc.add_paragraph(style="Title")
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.add_run("CampusFlow DSL Compiler Final Project Report")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Phase 3 Final Implementation and Testing")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = BLACK
    doc.add_paragraph()
    description = doc.add_paragraph()
    description.alignment = WD_ALIGN_PARAGRAPH.CENTER
    description.paragraph_format.left_indent = Inches(0.65)
    description.paragraph_format.right_indent = Inches(0.65)
    description.add_run(
        "Compiler Design Laboratory project covering lexical analysis, parsing, "
        "semantic validation, intermediate code, optimization, execution, and "
        "correctness-aware performance measurement."
    )
    doc.add_paragraph()
    meta = add_table(doc, ["Report item", "Value"], [
        ["Project", "CampusFlow domain-specific language compiler"],
        ["Submission", "Phase 3 final implementation and testing"],
        ["Implementation", "Python standard library, Python 3.10 or newer"],
        ["Verification date", date.today().isoformat()],
        ["Automated result", f"{facts['test_count']} tests passed"],
    ], [1.65, 4.55])
    doc.add_paragraph()
    purpose = doc.add_paragraph()
    purpose.alignment = WD_ALIGN_PARAGRAPH.CENTER
    purpose.add_run(
        "This report explains the delivered system, demonstrates its output, "
        "and maps the implementation to the Phase 3 review criteria."
    ).italic = True
    doc.add_page_break()

    doc.add_heading("Project Outcome", level=1)
    doc.add_paragraph(
        "CampusFlow is a working compiler and interpreter for a small event-"
        "scheduling language. The final pipeline accepts source text, produces "
        "positioned tokens and an abstract syntax tree, validates declarations "
        "and scheduling rules, generates linear intermediate code, removes "
        "provably unnecessary control flow, and executes the optimized program. "
        f"The verified suite contains {facts['test_count']} passing tests."
    )
    doc.add_paragraph(
        "The optimizer is conservative. It changes a condition only when the "
        "most recent reservation capacity proves the result. Benchmark mode "
        "executes both forms independently and compares their complete runtime "
        f"results. The recorded complex demonstration reports {facts['equivalence']} "
        "equivalence."
    )

    doc.add_heading("Contents", level=1)
    contents = [
        "1 Problem and Objective", "2 Language Definition", "3 Compiler Architecture",
        "4 Implementation by Module", "5 Optimization Algorithm", "6 Execution Engine",
        "7 Testing and Error Handling", "8 Performance Evaluation", "9 Demonstration Evidence",
        "10 User Instructions", "11 Review Criteria", "12 Limitations and Conclusion",
        "Appendix A Selected Source Listings",
    ]
    for item in contents:
        doc.add_paragraph(item)
    doc.add_page_break()

    doc.add_heading("1 Problem and Objective", level=1)
    doc.add_heading("1 1 Problem Definition", level=2)
    doc.add_paragraph(
        "Campus scheduling combines declarations, resource constraints, time "
        "conflicts, conditional actions, and cancellation. Encoding those rules "
        "in a small language creates a suitable compiler project because the "
        "same input must pass through lexical, syntactic, semantic, intermediate, "
        "optimization, and runtime stages."
    )
    doc.add_heading("1 2 Objective", level=2)
    doc.add_paragraph(
        "The objective is to implement and demonstrate the complete CampusFlow "
        "compiler. A successful run must reject invalid input at the correct "
        "stage, expose meaningful intermediate artifacts, execute valid schedules, "
        "and prove that optimization preserves observable behavior."
    )
    add_bullets(doc, [
        "Recognize keywords, identifiers, dates, times, numbers, operators, and delimiters.",
        "Build an AST with recursive-descent parsing and semicolon-based recovery.",
        "Validate declarations, cancellation, values, and schedule conflicts with a symbol table.",
        "Generate readable linear IR and safely fold capacity conditions.",
        "Execute optimized IR and report state, diagnostics, timing, memory, and equivalence.",
    ])

    doc.add_heading("2 Language Definition", level=1)
    doc.add_heading("2 1 Grammar", level=2)
    grammar = """program               = { statement } EOF ;
statement             = event_statement | reserve_statement
                      | assign_statement | cancel_statement | if_statement ;
event_statement       = "EVENT" IDENTIFIER "ON" DATE "AT" TIME "TO" TIME ";" ;
reserve_statement     = "RESERVE" IDENTIFIER "FOR" IDENTIFIER
                        "CAPACITY" NUMBER ";" ;
assign_statement      = "ASSIGN" IDENTIFIER "TO" IDENTIFIER ";" ;
cancel_statement      = "CANCEL" IDENTIFIER ";" ;
if_statement          = "IF" "CAPACITY" comparison NUMBER
                        "THEN" conditional_statement ;
comparison            = ">" | "<" | ">=" | "<=" | "==" | "!=" ;
conditional_statement = reserve_statement | assign_statement | cancel_statement ;"""
    add_code(doc, grammar, 8.2)
    doc.add_heading("2 2 Language Rules", level=2)
    add_table(doc, ["Construct", "Rule"], [
        ["EVENT", "Declares a unique, single-day event whose end is later than its start."],
        ["RESERVE", "Books a resource for an active event using a capacity greater than zero."],
        ["ASSIGN", "Attaches an item to an active declared event."],
        ["CANCEL", "Marks an event cancelled and releases current reservations and assignments."],
        ["IF capacity", "Compares the most recent successful reservation capacity and conditionally runs one statement."],
        ["Conflicts", "The same resource or item cannot serve overlapping active events; adjacent intervals are legal."],
    ], [1.25, 5.0])

    doc.add_heading("3 Compiler Architecture", level=1)
    doc.add_paragraph(
        "The pipeline is intentionally staged. Each stage returns a typed result "
        "and diagnostics. A failed stage prevents later stages from receiving "
        "unsafe or incomplete data."
    )
    architecture = (
        "Source  ->  Lexer  ->  Tokens  ->  Parser and AST  ->  Semantic Analysis\n"
        "        ->  Original IR  ->  Optimizer  ->  Optimized IR  ->  Interpreter"
    )
    add_code(doc, architecture, 10)
    add_table(doc, ["Stage", "Input", "Output", "Failure behavior"], [
        ["Lexer", "Characters", "Positioned tokens", "Reports lexical errors and stops parsing"],
        ["Parser", "Tokens", "AST", "Recovers at semicolons and records syntax errors"],
        ["Semantic", "AST", "Symbol table", "Rejects declaration, value, cancellation, and conflict errors"],
        ["IR generator", "Validated AST", "Original IR", "Runs only after semantic success"],
        ["Optimizer", "Original IR", "Validated optimized IR", "Returns original IR unchanged on structural errors"],
        ["Interpreter", "Optimized IR", "Runtime state and log", "Returns positioned runtime diagnostics"],
    ], [1.05, 1.15, 1.55, 2.5])

    doc.add_heading("4 Implementation by Module", level=1)
    add_table(doc, ["Module", "Implementation responsibility", "Primary tests"], [
        ["model.py", "Token, scan result, and lexical diagnostic records", "test_lexer.py"],
        ["lexer.py", "Longest-match scanning, comments, positions, date/time validation", "test_lexer.py"],
        ["ast_nodes.py", "Frozen AST records and deterministic tree formatter", "test_ast_nodes.py"],
        ["parser.py", "Recursive-descent grammar and statement-level recovery", "test_parser.py"],
        ["symbols.py", "Events, reservations, assignments, and interval overlap", "test_symbols.py"],
        ["semantic.py", "Source-ordered declarations, values, cancellation, and conflict checks", "test_semantic.py"],
        ["ir.py", "Instruction model, lowering, labels, and stable formatting", "test_ir.py"],
        ["optimizer.py", "IR validation, capacity propagation, folding, and statistics", "test_optimizer.py"],
        ["interpreter.py", "Defensive IR execution and deterministic runtime state", "test_interpreter.py"],
        ["metrics.py", "Warm-up, median time, peak memory, and full-result equivalence", "test_metrics.py"],
        ["pipeline.py", "Stage gating and optimized execution", "test_pipeline.py"],
        ["cli.py", "Arguments, stages, diagnostics, output, and benchmark mode", "test_cli.py"],
    ], [1.15, 3.7, 1.35])

    doc.add_heading("5 Optimization Algorithm", level=1)
    doc.add_heading("5 1 Forward Capacity Propagation", level=2)
    doc.add_paragraph(
        "The optimizer scans IR from left to right and stores the capacity from "
        "the latest unconditional reservation. It first builds and validates the "
        "label table. Duplicate labels, missing targets, backward conditional "
        "jumps, malformed operands, and unknown operators become positioned "
        "optimizer diagnostics."
    )
    add_bullets(doc, [
        "Known true condition: remove the conditional jump and end label, retain the body, and update known capacity if the body reserves a resource.",
        "Known false condition: remove the jump, unreachable body, and end label.",
        "Unknown condition: preserve the complete region and invalidate capacity knowledge so later folds remain safe.",
        "Nested or unfamiliar flow: preserve the region because the pass cannot prove one resulting capacity.",
    ])
    doc.add_heading("5 2 Before and After Example", level=2)
    before = """0000  CREATE_EVENT 'Demo', '2026-10-15', '10:00', '11:00'
0001  RESERVE_RESOURCE 'MainHall', 'Demo', 100
0002  JUMP_IF_FALSE '>=', 100, 'if_end_1'
0003  ASSIGN_ITEM 'Projector', 'Demo'
0004  LABEL 'if_end_1'
0005  JUMP_IF_FALSE '<', 50, 'if_end_2'
0006  ASSIGN_ITEM 'BackupSpeaker', 'Demo'
0007  LABEL 'if_end_2'"""
    after = """0000  CREATE_EVENT 'Demo', '2026-10-15', '10:00', '11:00'
0001  RESERVE_RESOURCE 'MainHall', 'Demo', 100
0002  ASSIGN_ITEM 'Projector', 'Demo'"""
    doc.add_paragraph("Original intermediate code")
    add_code(doc, before, 8)
    doc.add_paragraph("Optimized intermediate code")
    add_code(doc, after, 8)
    doc.add_paragraph(
        f"The recorded example decreases from {facts['demo_original']} to "
        f"{facts['demo_optimized']} instructions, a {facts['demo_reduction']} "
        "reduction. It folds one true condition and one false condition."
    )

    doc.add_page_break()
    doc.add_heading("6 Execution Engine", level=1)
    doc.add_paragraph(
        "The interpreter validates every opcode and operand before applying it. "
        "It stores events in insertion order and maintains reservation and "
        "assignment lists. Cancellation removes associated runtime records, "
        "which allows a later event to reuse released resources. Conditional "
        "jumps compare the last successful reservation capacity and can only "
        "move forward to a validated label."
    )
    add_table(doc, ["Runtime record", "Stored information", "Main invariant"], [
        ["Event", "Name, date, start, end, cancelled flag", "Unique name and valid same-day time range"],
        ["Reservation", "Resource, event, capacity", "Positive capacity and no overlapping resource use"],
        ["Assignment", "Item and event", "No overlapping item use"],
        ["Execution result", "Final state and diagnostics", "Deterministic output; first defensive error stops execution"],
    ], [1.35, 2.3, 2.55])

    doc.add_heading("7 Testing and Error Handling", level=1)
    doc.add_heading("7 1 Test Matrix", level=2)
    add_table(doc, ["Area", "Coverage", "Representative evidence"], [
        ["Lexical", "Tokens, longest operators, comments, positions, illegal input, date/time limits", "test_lexer.py and invalid.cflow"],
        ["Syntax and AST", "All statements, recovery, missing delimiters, case-insensitive keywords, immutable nodes", "test_parser.py and test_ast_nodes.py"],
        ["Semantic", "Duplicates, undeclared/cancelled names, capacities, overlaps, conditional effects", "test_semantic.py and resource_conflict.cflow"],
        ["IR and runtime", "Lowering, formatting, execution, malformed opcodes, missing capacity, backward jumps", "test_ir.py and test_interpreter.py"],
        ["Optimization", "True/false folding, state propagation, unknown facts, malformed control flow", "test_optimizer.py"],
        ["Metrics", "Iteration validation, equivalence failure, timing noise, stable report", "test_metrics.py"],
        ["End to end", "CLI stages, failures, boundaries, complex schedule, benchmark", "test_cli.py and test_phase3_examples.py"],
    ], [1.1, 3.25, 1.85])
    doc.add_heading("7 2 Verified Result", level=2)
    doc.add_paragraph(
        f"The command `python3 -m unittest discover -s tests -v` completed "
        f"successfully with {facts['test_count']} tests. The evidence file records "
        "the full test names and final OK status."
    )
    doc.add_heading("7 3 Error Demonstrations", level=2)
    add_table(doc, ["Error stage", "Demonstration", "Expected response"], [
        ["Lexical", "invalid.cflow", "Illegal characters, invalid date/time, operator, or string with positions"],
        ["Syntax", "syntax_errors.cflow", "Unexpected or missing grammar element; recovery at semicolon"],
        ["Semantic", "semantic_errors.cflow", "Duplicate, undeclared, cancelled, capacity, or scheduling diagnostic"],
        ["Optimizer", "Malformed IR unit tests", "Control-flow diagnostic and original IR retained"],
        ["Runtime", "Malformed IR unit tests", "Defensive opcode, operand, label, or state diagnostic"],
    ], [1.1, 1.65, 3.45])

    doc.add_heading("8 Performance Evaluation", level=1)
    doc.add_heading("8 1 Method", level=2)
    doc.add_paragraph(
        "Benchmark mode warms each program once, then executes every measured "
        "iteration with a fresh interpreter. It records integer median time with "
        "a nanosecond timer and measures peak traced memory separately. The "
        "complete original and optimized execution results must compare equal."
    )
    doc.add_heading("8 2 Recorded Complex Schedule Result", level=2)
    add_table(doc, ["Metric", "Original", "Optimized"], [
        ["Instruction count", str(facts["bench_original"]), str(facts["bench_optimized"])],
        ["Median execution time in ns", str(facts["original_ns"]), str(facts["optimized_ns"])],
        ["Peak traced memory in bytes", str(facts["original_bytes"]), str(facts["optimized_bytes"])],
    ], [2.75, 1.7, 1.7])
    doc.add_paragraph(
        f"The recorded run used {facts['bench_iterations']} iterations on "
        f"{facts['python']} and {facts['platform']}. Instruction count fell by "
        f"{facts['bench_reduction']}. The measured original-to-optimized time "
        f"ratio was {facts['ratio']}, and execution equivalence was "
        f"{facts['equivalence']}. Timing is environment-dependent; equivalence "
        "and instruction reduction are the primary correctness checks."
    )

    doc.add_heading("9 Demonstration Evidence", level=1)
    doc.add_heading("9 1 Boundary Demonstration", level=2)
    doc.add_paragraph(
        "boundary_cases.cflow uses capacity 1, equality and inequality operators, "
        "and two events meeting at exactly 11:00. Both events legally reserve the "
        "same room because the intervals are adjacent rather than overlapping."
    )
    doc.add_heading("9 2 Complex Demonstration", level=2)
    doc.add_paragraph(
        "complex_schedule.cflow declares six events and exercises multiple "
        "resources, assignments, true and false conditions, cancellation, and "
        "legal resource reuse. The optimizer removes ten of twenty-nine "
        "instructions. Execution cancels Opening, retains active later events, "
        "and reuses MainHall for Workshop after the cancellation."
    )
    complex_excerpt = """OPTIMIZATION SUMMARY
Original instructions: 29
Optimized instructions: 19
Instructions removed: 10
Instruction reduction: 34.48%
Conditions folded true: 2
Conditions folded false: 2

FINAL RUNTIME STATE
Opening [CANCELLED] 2026-10-25 09:00-10:30
Workshop [ACTIVE] 2026-10-25 10:00-12:00
  RESERVATION LabA capacity=80
  RESERVATION MainHall capacity=100
  ASSIGNMENT Projector
Closing [ACTIVE] 2026-10-25 16:00-17:00
  RESERVATION MainHall capacity=200
  ASSIGNMENT StageLights"""
    add_code(doc, complex_excerpt, 8)

    doc.add_heading("10 User Instructions", level=1)
    doc.add_heading("10 1 Run the Project", level=2)
    doc.add_paragraph("macOS or Linux from the repository root")
    add_code(doc, "cd phase3\npython3 run.py examples/complex_schedule.cflow")
    doc.add_paragraph("Windows PowerShell from the repository root")
    add_code(doc, "cd phase3\npy run.py examples\\complex_schedule.cflow")
    doc.add_heading("10 2 Inspect One Compiler Stage", level=2)
    add_code(doc, "python3 run.py examples/optimization_demo.cflow --stage optimized")
    doc.add_paragraph(
        "Available stages are tokens, ast, ir, optimized, run, and all. The "
        "default is all. Successful runs return status 0, compiler/runtime errors "
        "return 1, and command-line or file errors return 2."
    )
    doc.add_heading("10 3 Benchmark and Test", level=2)
    add_code(doc, "python3 run.py examples/complex_schedule.cflow --benchmark 100\npython3 -m unittest discover -s tests -v")

    doc.add_heading("11 Review Criteria", level=1)
    add_table(doc, ["Criterion", "Delivered evidence"], [
        ["Project Completion", "Complete staged compiler, optimizer, interpreter, CLI, metrics, examples, tests, report, and package"],
        ["Functional Correctness", f"{facts['test_count']} tests pass and benchmark equivalence is {facts['equivalence']}"],
        ["Technical Complexity", "Recursive descent, symbol validation, IR lowering, forward data-flow folding, runtime, and metrics"],
        ["Compiler Integration", "One gated pipeline executes optimized IR only after every prior stage succeeds"],
        ["Innovation", "Campus scheduling DSL with conservative capacity folding and correctness-aware benchmarking"],
        ["Code Quality", "Typed dataclasses, small modules, deterministic formatters, source positions, standard library"],
        ["Testing", "Unit, integration, CLI, invalid, boundary, malformed IR, noise, and complex scenario tests"],
        ["Demonstration", "Optimization, boundary, complex, and benchmark examples with captured output"],
        ["Documentation", "User guide, rubric map, evidence, final report, presentation, and source listings"],
        ["Viva", "This report explains the grammar, architecture, algorithm, commands, evidence, and limitations"],
    ], [1.35, 4.85])

    doc.add_heading("12 Limitations and Conclusion", level=1)
    doc.add_heading("12 1 Limitations", level=2)
    add_bullets(doc, [
        "Runtime state is in memory and is not connected to an external database or calendar.",
        "Events are single-day; an event cannot cross midnight.",
        "The only conditional value is the most recent reservation capacity.",
        "An IF body contains one non-EVENT statement; the source grammar has no nested IF, loops, functions, or user variables.",
        "The project interprets optimized IR and does not emit native machine code.",
        "Benchmark time and memory values depend on the machine and current system load.",
    ])
    doc.add_heading("12 2 Conclusion", level=2)
    doc.add_paragraph(
        "CampusFlow meets the Phase 3 objective as a complete, demonstrable "
        "compiler project. The staged implementation applies compiler concepts "
        "from tokenization through optimized execution. Its optimizer reduces "
        "provable control flow while preserving uncertain regions, and the "
        "benchmark confirms complete execution-result equivalence. The automated "
        "suite, recorded outputs, user instructions, and selected source listings "
        "provide reproducible evidence for the final review."
    )

    doc.add_page_break()
    doc.add_heading("Appendix A Selected Source Listings", level=1)
    doc.add_paragraph(
        "The following excerpts are read directly from the final source tree when "
        "this report is generated. They show the central optimization, pipeline, "
        "and benchmark implementations."
    )
    doc.add_heading("A 1 Optimizer Pass", level=2)
    optimizer_code = extract_block(
        PHASE3 / "campusflow" / "optimizer.py",
        "def optimize",
        "@staticmethod",
        58,
    )
    add_code(doc, optimizer_code, 7.2)
    doc.add_page_break()
    doc.add_heading("A 2 Compiler Pipeline", level=2)
    pipeline_code = extract_block(
        PHASE3 / "campusflow" / "pipeline.py",
        "def compile_source",
        "execution = Interpreter",
        54,
    )
    add_code(doc, pipeline_code, 7.2)
    doc.add_page_break()
    doc.add_heading("A 3 Benchmark Function", level=2)
    metrics_code = extract_block(
        PHASE3 / "campusflow" / "metrics.py",
        "def benchmark_ir",
        "def _time_execution",
        48,
    )
    add_code(doc, metrics_code, 7.2)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build_report(require_tree())
